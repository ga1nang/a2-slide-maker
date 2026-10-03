#!/usr/bin/env python3
"""Exportiert die Köln-Präsentation (index.html) als PowerPoint (.pptx).

Jede Folie wird mit headless Chrome als Bild gerendert und danach randlos in
eine PowerPoint-Folie eingebettet. Das Layout (Karten, Sticker, Schriften,
Farbverläufe, Rotation) ist reines CSS und lässt sich in PowerPoint nicht
nachbauen – deshalb steckt hier bewusst ein Bild pro Folie, kein Nachbau mit
PowerPoint-Formen. Der Text ist dadurch nicht mehr editierbar, sieht aber
exakt so aus wie im Browser.

Benutzung:
    python3 export_pptx.py                  # -> koeln-praesentation.pptx
    python3 export_pptx.py -o backup.pptx
    python3 export_pptx.py --open            # danach direkt öffnen
    python3 export_pptx.py --width 2560      # höhere Auflösung (Standard: 1920)

Braucht die Pakete "python-pptx" und "Pillow". Falls sie fehlen, legt das
Skript beim ersten Aufruf automatisch eine eigene virtuelle Umgebung an
(`.venv-pptx/`) und startet sich darin neu – kein manuelles `pip install`
nötig, auch wenn das System-Python "externally managed" ist (PEP 668).
"""

from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent


def _ensure_dependencies() -> None:
    """Sorgt dafür, dass python-pptx/Pillow importierbar sind – notfalls über
    eine selbst angelegte virtuelle Umgebung, in der das Skript neu startet."""
    try:
        import pptx  # noqa: F401
        import PIL  # noqa: F401
        return
    except ImportError:
        pass

    venv_dir = HERE / ".venv-pptx"
    venv_python = venv_dir / ("Scripts/python.exe" if os.name == "nt" else "bin/python3")

    if not venv_python.exists():
        print("→ Richte einmalig eine kleine Python-Umgebung für den PPTX-Export ein …")
        subprocess.run(
            [sys.executable, "-m", "venv", "--system-site-packages", str(venv_dir)],
            check=True,
        )
        subprocess.run(
            [str(venv_python), "-m", "pip", "install", "--quiet", "--upgrade", "pip"],
            check=True,
        )
        subprocess.run(
            [str(venv_python), "-m", "pip", "install", "--quiet", "python-pptx", "Pillow"],
            check=True,
        )
        print("→ Fertig eingerichtet.\n")

    os.execv(str(venv_python), [str(venv_python), str(Path(__file__).resolve()), *sys.argv[1:]])


_ensure_dependencies()

import argparse  # noqa: E402
import re  # noqa: E402
import shutil  # noqa: E402
import tempfile  # noqa: E402

from PIL import Image, ImageStat  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

sys.path.insert(0, str(HERE))
from export_pdf import BROWSERS, find_browser  # noqa: E402

# Standard-Breitbildformat, das PowerPoint selbst für 16:9 verwendet.
SLIDE_WIDTH_EMU = 12192000   # 13.333 in
SLIDE_HEIGHT_EMU = 6858000   # 7.5 in

# Chrome reserviert im Headless-Modus konstant ~87px Höhe, die nicht zum
# sichtbaren Viewport gehören – unabhängig von der gewählten Fenstergröße
# (siehe CLAUDE.md unter "Commands"; empirisch bei 720/1080/1440px geprüft).
CHROME_CHROME_PX = 87


def count_slides(html: Path) -> int:
    text = html.read_text(encoding="utf-8")
    return len(re.findall(r'<section class="slide\b', text))


# Unterhalb dieser Standardabweichung (Graustufen) gilt eine Folie als
# vermutlich leer fotografiert – kalibriert an einem echten Fehlerfall
# (kaputte Folie: ~8, jede echte Folie im Deck: >30).
BLANK_STDDEV_THRESHOLD = 20


def _is_probably_blank(img: Image.Image) -> bool:
    return ImageStat.Stat(img.convert("L")).stddev[0] < BLANK_STDDEV_THRESHOLD


def screenshot_slide(
    browser: str, html_uri: str, index: int, width: int, height: int,
    budget: int, timeout: int, out_path: Path,
) -> None:
    # slides.js schaltet bei ?static alle Animationen sofort ab. Trotzdem als
    # zweite Absicherung: Chromes --virtual-time-budget garantiert nicht,
    # dass der Compositor fertig gemalt hat, bevor der Screenshot ausgelöst
    # wird – bei einem Testlauf wurde so einmal eine Folie halbtransparent
    # eingefangen. Deshalb nach dem Rendern kurz prüfen und notfalls einmal
    # wiederholen, statt das stillschweigend in die PPTX zu übernehmen.
    for attempt in (1, 2):
        with tempfile.TemporaryDirectory(prefix="chrome-pptx-") as profile:
            cmd = [
                browser,
                "--headless=new",
                "--disable-gpu",
                "--no-sandbox",
                "--no-first-run",
                f"--user-data-dir={profile}",
                "--hide-scrollbars",
                "--force-device-scale-factor=1",
                f"--virtual-time-budget={budget}",
                f"--window-size={width},{height + CHROME_CHROME_PX}",
                f"--screenshot={out_path}",
                f"{html_uri}?static#{index}",
            ]
            try:
                proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
            except subprocess.TimeoutExpired:
                sys.exit(f"Abbruch: Folie {index} hat länger als {timeout}s gebraucht.")

        if not out_path.is_file() or out_path.stat().st_size == 0:
            sys.stderr.write(proc.stderr[-1500:] + "\n")
            sys.exit(f"Screenshot von Folie {index} fehlgeschlagen (Exit-Code {proc.returncode}).")

        # Chrome liefert width x (height+87) – auf die exakte 16:9-Bühne zuschneiden.
        img = Image.open(out_path)
        img = img.crop((0, 0, width, height))
        img.save(out_path)

        if not _is_probably_blank(img):
            return
        if attempt == 1:
            print(f"  Folie {index}: sieht fast leer aus, versuche es noch einmal …")

    sys.exit(
        f"Folie {index} wurde zweimal fast leer fotografiert. "
        f"Größeres --budget versuchen (aktuell {budget}ms)."
    )


def build_pptx(image_paths: list[Path], out: Path) -> None:
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_WIDTH_EMU)
    prs.slide_height = Emu(SLIDE_HEIGHT_EMU)
    blank_layout = prs.slide_layouts[6]  # leeres Layout, keine Platzhalter

    for img_path in image_paths:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            str(img_path), 0, 0, width=prs.slide_width, height=prs.slide_height
        )

    prs.save(out)


def export(
    html: Path, out: Path, browser: str, width: int, budget: int, timeout: int
) -> None:
    if not html.is_file():
        sys.exit(f"Datei nicht gefunden: {html}")

    n = count_slides(html)
    if n == 0:
        sys.exit(f"Keine Folien in {html} gefunden.")

    height = round(width * 9 / 16)
    out.parent.mkdir(parents=True, exist_ok=True)

    print(f"→ Browser    : {browser}")
    print(f"→ Quelle     : {html}")
    print(f"→ Auflösung  : {width}x{height} px pro Folie, {n} Folien")
    print(f"→ Ziel       : {out}")

    with tempfile.TemporaryDirectory(prefix="koeln-pptx-") as tmp:
        tmp_dir = Path(tmp)
        images = []
        for i in range(1, n + 1):
            img_path = tmp_dir / f"slide-{i:02d}.png"
            screenshot_slide(browser, html.as_uri(), i, width, height, budget, timeout, img_path)
            images.append(img_path)
            print(f"  Folie {i:>2}/{n} gerendert")

        build_pptx(images, out)

    size_mb = out.stat().st_size / 1024 / 1024
    print(f"✓ Fertig: {n} Folien, {size_mb:.1f} MB")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Köln-Präsentation als PowerPoint exportieren (ein Bild pro Folie)."
    )
    parser.add_argument(
        "html", nargs="?", type=Path, default=HERE / "index.html",
        help="HTML-Datei der Präsentation (Standard: index.html neben diesem Skript)",
    )
    parser.add_argument(
        "-o", "--out", type=Path, default=HERE / "koeln-praesentation.pptx",
        help="Zieldatei (Standard: koeln-praesentation.pptx)",
    )
    parser.add_argument(
        "--browser", help="Pfad/Name des Browsers, falls er nicht automatisch gefunden wird"
    )
    parser.add_argument(
        "--width", type=int, default=1920,
        help="Bildbreite pro Folie in Pixeln, Höhe folgt aus 16:9 (Standard: 1920)",
    )
    parser.add_argument(
        "--budget", type=int, default=8000,
        help="Wartezeit in ms pro Folie für Schriften/Bilder (Standard: 8000)",
    )
    parser.add_argument(
        "--timeout", type=int, default=60,
        help="Maximale Laufzeit pro Folie in Sekunden (Standard: 60)",
    )
    parser.add_argument("--open", action="store_true", help="PPTX nach dem Export öffnen")
    args = parser.parse_args()

    export(
        html=args.html.resolve(),
        out=args.out.resolve(),
        browser=find_browser(args.browser),
        width=args.width,
        budget=args.budget,
        timeout=args.timeout,
    )

    if args.open:
        opener = shutil.which("xdg-open") or shutil.which("open")
        if opener:
            subprocess.Popen([opener, str(args.out.resolve())])


if __name__ == "__main__":
    main()
